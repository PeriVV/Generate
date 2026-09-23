#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""从 docs/java智能单元测试工具使用手册.md 生成排版规范的中国文档风格 .docx。

用法：python3 docs/build_manual_docx.py
样式：黑体标题、宋体正文、五号表格、带边框表格、封面、目录域、页脚页码。
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

SRC = Path(__file__).with_name("java智能单元测试工具使用手册.md")
DST = Path(__file__).with_name("java智能单元测试工具使用手册.docx")

HEI = "黑体"
SONG = "宋体"
LATIN = "Times New Roman"


def set_east_asia(style_or_run, east: str, latin: str) -> None:
    """同时设置中西文字体，否则中文字体不生效。"""
    rpr = style_or_run.element.get_or_add_rPr() if hasattr(style_or_run, "element") else style_or_run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east)


def add_run(p, text: str, *, font=SONG, latin=LATIN, size=12, bold=False, color=None):
    run = p.add_run(text)
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    set_east_asia(run, font, latin)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if color:
        run.font.color.rgb = color
    return run


def body_para(doc, text: str, *, size=12, font=SONG, bold=False, align=None,
              space_after=6, line=1.5, indent=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    if line:
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = line
    if align is not None:
        p.alignment = align
    if indent:
        pf.first_line_indent = Pt(size * 2)
    add_run(p, text, font=font, size=size, bold=bold)
    return p


def heading(doc, text: str, level: int):
    p = doc.add_paragraph(style=f"Heading {level}")
    sizes = {1: 16, 2: 14, 3: 12}
    size = sizes[level]
    pf = p.paragraph_format
    pf.space_before = Pt({1: 18, 2: 12, 3: 8}[level])
    pf.space_after = Pt({1: 12, 2: 8, 3: 6}[level])
    pf.line_spacing = 1.3
    pf.keep_with_next = True
    add_run(p, text, font=HEI, size=size, bold=True, color=RGBColor(0, 0, 0))
    return p


def shade_cell(cell, fill: str):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(r'<w:shd {} w:val="clear" w:color="auto" w:fill="{}"/>'.format(nsdecls("w"), fill)))


def add_table(doc, rows, widths_equal=False):
    """带边框表格，表头加底色加粗，五号宋体。"""
    n_cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.vertical_alignment = 1  # center
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if r == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(2)
            add_run(p, text, size=10.5, bold=(r == 0))
            if r == 0:
                shade_cell(cell, "E8E8E8")
            if widths_equal:
                cell.width = Cm(16.0 / n_cols)
    return table


def figure_box(doc, hint: str):
    """截图占位框：带边框单元格，高约 6cm，灰字居中。"""
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    row = table.rows[0]
    row.height = Cm(6)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    cell = table.cell(0, 0)
    cell.width = Cm(15.5)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(72)
    p.paragraph_format.space_after = Pt(72)
    add_run(p, hint, size=10.5, color=RGBColor(0x80, 0x80, 0x80))
    return table


def add_field(p, instr: str, hint: str = ""):
    """插入 Word 域（目录 / 页码）。"""
    r1 = p.add_run()
    r1._element.append(OxmlElement("w:fldChar"))
    r1._element[-1].set(qn("w:fldCharType"), "begin")
    r2 = p.add_run()
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    r2._element.append(instr_el)
    r3 = p.add_run()
    r3._element.append(OxmlElement("w:fldChar"))
    r3._element[-1].set(qn("w:fldCharType"), "separate")
    if hint:
        add_run(p, hint, size=10.5, color=RGBColor(0x80, 0x80, 0x80))
    r4 = p.add_run()
    r4._element.append(OxmlElement("w:fldChar"))
    r4._element[-1].set(qn("w:fldCharType"), "end")


def build():
    doc = Document()
    # 全局样式
    normal = doc.styles["Normal"]
    normal.font.name = LATIN
    normal.font.size = Pt(12)
    set_east_asia(normal, SONG, LATIN)
    for sec in doc.sections:
        sec.top_margin = Cm(2.54)
        sec.bottom_margin = Cm(2.54)
        sec.left_margin = Cm(3.0)
        sec.right_margin = Cm(3.0)
        # 页脚页码：第 X 页 共 Y 页
        footer_p = sec.footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(footer_p, "第 ", size=9)
        add_field(footer_p, "PAGE")
        add_run(footer_p, " 页 共 ", size=9)
        add_field(footer_p, "NUMPAGES")
        add_run(footer_p, " 页", size=9)
        for run in footer_p.runs:
            run.font.size = Pt(9)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    toc_inserted = False
    first_h1_seen = False

    while i < len(lines):
        line = lines[i].rstrip()

        # 空行
        if not line.strip():
            i += 1
            continue

        # 标题
        if line.startswith("### "):
            heading(doc, line[4:].strip(), 3)
            i += 1
            continue
        if line.startswith("## "):
            heading(doc, line[3:].strip(), 2)
            i += 1
            continue
        if line.startswith("# "):
            if not first_h1_seen:
                first_h1_seen = True
                # 封面结束后：目录页
                doc.add_page_break()
                body_para(doc, "目  录", font=HEI, size=16, bold=True,
                          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
                toc_p = doc.add_paragraph()
                add_field(toc_p, r'TOC \o "1-2" \h \z \u',
                          hint="（在 Word / WPS 中右键此处目录，选择“更新域”即可生成章节页码）")
                doc.add_page_break()
                toc_inserted = True
            heading(doc, line[2:].strip(), 1)
            i += 1
            continue

        # 表格块
        if line.lstrip().startswith("|"):
            block = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                raw = lines[i].strip().strip("|")
                cells = [c.strip() for c in raw.split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    block.append(cells)
                i += 1
            add_table(doc, block, widths_equal=("更改单编号" in block[0]))
            continue

        # 截图占位（加粗的【待插入截图…】）
        m = re.fullmatch(r"\*\*(【待插入截图：.*?】)\*\*", line.strip())
        if m:
            figure_box(doc, m.group(1))
            i += 1
            continue

        # 表题 / 图题
        if re.match(r"^(表|图) \d+-\d+ ", line.strip()):
            body_para(doc, line.strip(), font=HEI, size=10.5, bold=True,
                      align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4, line=1.2)
            i += 1
            continue

        # 全行加粗段（封面要素 / 签署页）
        m = re.fullmatch(r"\*\*(.+?)\*\*", line.strip())
        if m:
            text = m.group(1)
            if text.startswith("（此处可插入封面"):
                i += 1
                continue
            if text == "java智能单元测试工具使用手册":
                for _ in range(6):
                    body_para(doc, "", space_after=6)
                body_para(doc, text, font=HEI, size=22, bold=True,
                          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)
                body_para(doc, "", space_after=6)
            elif text == "签署页":
                doc.add_page_break()
                body_para(doc, "签署页", font=HEI, size=16, bold=True,
                          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
            else:
                body_para(doc, text, font=HEI, size=14, bold=True, space_after=8)
            i += 1
            continue

        # 正文段（去掉行内 ** 与 ` 标记）
        text = line.strip().replace("**", "").replace("`", "")
        body_para(doc, text)
        i += 1

    doc.core_properties.title = "java智能单元测试工具使用手册"
    doc.save(DST)
    print(f"生成完成：{DST}")


if __name__ == "__main__":
    build()
